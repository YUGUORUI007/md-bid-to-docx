from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION

from common import TODO_RE, write_json, write_text


MERMAID_RE = re.compile(r"```+\s*mermaid|flowchart\s+(TD|TB|LR|RL)|sequenceDiagram|stateDiagram-v2", re.IGNORECASE)
INTERNAL_PROMPT_RE = re.compile(r"自动排版生成|请按需补充|右键更新域|Markdown 自动|由 Markdown", re.IGNORECASE)


def audit_docx(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_texts: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    table_texts.append(cell.text)
    all_text = "\n".join(paragraphs + table_texts)
    placeholders = []
    for match in TODO_RE.finditer(all_text):
        start = max(match.start() - 40, 0)
        end = min(match.end() + 100, len(all_text))
        placeholders.append(all_text[start:end].replace("\n", " ")[:180])
    mermaid_leftovers = []
    for match in MERMAID_RE.finditer(all_text):
        start = max(match.start() - 40, 0)
        end = min(match.end() + 100, len(all_text))
        mermaid_leftovers.append(all_text[start:end].replace("\n", " ")[:180])
    internal_prompts = []
    for match in INTERNAL_PROMPT_RE.finditer(all_text):
        start = max(match.start() - 40, 0)
        end = min(match.end() + 100, len(all_text))
        internal_prompts.append(all_text[start:end].replace("\n", " ")[:180])

    image_count = 0
    with zipfile.ZipFile(path) as zf:
        image_count = sum(1 for name in zf.namelist() if name.startswith("word/media/"))
        has_document = "word/document.xml" in zf.namelist()
        has_styles = "word/styles.xml" in zf.namelist()

    section_count = len(doc.sections)
    landscape_section_count = sum(
        1 for section in doc.sections if section.orientation == WD_ORIENTATION.LANDSCAPE or section.page_width > section.page_height
    )

    issues = []
    if not has_document:
        issues.append({"severity": "fatal", "kind": "missing-document-xml", "message": "DOCX has no word/document.xml."})
    if not has_styles:
        issues.append({"severity": "warning", "kind": "missing-styles", "message": "DOCX has no word/styles.xml."})
    if mermaid_leftovers:
        issues.append({"severity": "fatal", "kind": "mermaid-leftover", "message": f"{len(mermaid_leftovers)} Mermaid/code leftovers detected."})
    if internal_prompts:
        issues.append({"severity": "warning", "kind": "internal-prompts-left", "message": f"{len(internal_prompts)} internal prompt/update markers detected."})
    if placeholders:
        issues.append({"severity": "warning", "kind": "placeholders-left", "message": f"{len(placeholders)} placeholder markers detected."})
    if len(paragraphs) < 3:
        issues.append({"severity": "warning", "kind": "low-paragraph-count", "message": "DOCX has very few paragraphs."})

    return {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
        "image_count": image_count,
        "section_count": section_count,
        "landscape_section_count": landscape_section_count,
        "placeholder_count": len(placeholders),
        "mermaid_leftover_count": len(mermaid_leftovers),
        "internal_prompt_count": len(internal_prompts),
        "placeholder_samples": placeholders[:20],
        "mermaid_leftover_samples": mermaid_leftovers[:20],
        "internal_prompt_samples": internal_prompts[:20],
        "fatal_count": sum(1 for issue in issues if issue["severity"] == "fatal"),
        "warning_count": sum(1 for issue in issues if issue["severity"] == "warning"),
        "issues": issues,
    }


def audit_to_markdown(report: dict) -> str:
    lines = [
        "# DOCX Structural Audit",
        "",
        f"- File: `{report['path']}`",
        f"- Size: {report['size_bytes']} bytes",
        f"- Paragraphs: {report['paragraph_count']}",
        f"- Tables: {report['table_count']}",
        f"- Images: {report['image_count']}",
        f"- Sections: {report['section_count']}",
        f"- Landscape sections: {report['landscape_section_count']}",
        f"- Placeholders: {report['placeholder_count']}",
        f"- Mermaid leftovers: {report['mermaid_leftover_count']}",
        f"- Internal prompt markers: {report['internal_prompt_count']}",
        f"- Fatal issues: {report['fatal_count']}",
        f"- Warnings: {report['warning_count']}",
        "",
        "## Issues",
        "",
    ]
    if not report["issues"]:
        lines.append("No structural issues detected.")
    else:
        lines.append("| Severity | Kind | Message |")
        lines.append("| --- | --- | --- |")
        for issue in report["issues"]:
            lines.append(f"| {issue['severity']} | {issue['kind']} | {issue['message']} |")
    if report["placeholder_samples"]:
        lines.extend(["", "## Placeholder Samples", ""])
        for sample in report["placeholder_samples"]:
            lines.append(f"- {sample}")
    if report["mermaid_leftover_samples"]:
        lines.extend(["", "## Mermaid Leftover Samples", ""])
        for sample in report["mermaid_leftover_samples"]:
            lines.append(f"- {sample}")
    if report["internal_prompt_samples"]:
        lines.extend(["", "## Internal Prompt Samples", ""])
        for sample in report["internal_prompt_samples"]:
            lines.append(f"- {sample}")
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit a DOCX structurally.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--out-dir", type=Path, default=None)
    args = parser.parse_args()
    input_path = args.input.resolve()
    if not input_path.is_file():
        parser.error(f"DOCX not found: {input_path}")
    out_dir = (args.out_dir or input_path.parent).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    report = audit_docx(input_path)
    write_json(out_dir / "docx-audit.json", report)
    write_text(out_dir / "docx-audit.md", audit_to_markdown(report))
    print(f"Wrote {out_dir / 'docx-audit.md'}")
    return 0 if report["fatal_count"] == 0 else 6


if __name__ == "__main__":
    raise SystemExit(main())
