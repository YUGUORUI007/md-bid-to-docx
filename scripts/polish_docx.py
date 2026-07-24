from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from common import DEFAULT_COMPANY_NAME, DEFAULT_COMPANY_SLOGAN, DEFAULT_COMPANY_WEBSITE, style_preset


BODY_FIRST_LINE_INDENT = Cm(0.74)
LIST_TEXT_INDENT = Cm(0.74)
LIST_HANGING_INDENT = Cm(-0.74)
TABLE_NOTE_RE = re.compile(r"^表\d+(?:[-－—]\d+[A-Za-z]?)?\s+.+$")
CN_SECTION_TITLE_RE = re.compile(r"^[（(][一二三四五六七八九十0-9]+[）)]")


def set_run_font(run, east_asia: str, ascii_font: str, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = ascii_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    run._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    paragraph.add_run("")
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_page_field(paragraph) -> None:
    paragraph.add_run("第")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" 页")


def add_paragraph_border(paragraph, edge: str, color: str, size: str = "8", space: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = p_bdr.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        p_bdr.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)


def clear_story(story):
    for child in list(story._element):
        story._element.remove(child)
    return story.add_paragraph()


def reset_paragraph_rhythm(paragraph) -> None:
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def section_text_width(section):
    return section.page_width - section.left_margin - section.right_margin


def apply_branded_page_furniture(
    section,
    preset: dict,
    title: str,
    *,
    blank_first_page: bool = False,
    company_name: str = DEFAULT_COMPANY_NAME,
    website: str = DEFAULT_COMPANY_WEBSITE,
    slogan: str = DEFAULT_COMPANY_SLOGAN,
) -> None:
    text_width = section_text_width(section)
    section.header_distance = Cm(0.85)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = blank_first_page
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    if blank_first_page:
        clear_story(section.first_page_header)
        clear_story(section.first_page_footer)

    header = clear_story(section.header)
    reset_paragraph_rhythm(header)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    left = header.add_run(company_name)
    set_run_font(left, preset["heading_font_east_asia"], preset["heading_font_ascii"], 8.5, True, preset["accent"])
    header.add_run("\t")
    right = header.add_run(title)
    set_run_font(right, preset["font_east_asia"], preset["font_ascii"], 8, color="666666")
    add_paragraph_border(header, "bottom", preset["accent"], size="6", space="2")

    footer = clear_story(section.footer)
    reset_paragraph_rhythm(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.paragraph_format.tab_stops.add_tab_stop(text_width // 2, WD_TAB_ALIGNMENT.CENTER)
    footer.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    slogan_display = " / ".join(str(slogan).split()) or str(slogan)
    footer.add_run(slogan_display)
    footer.add_run("\t")
    add_page_field(footer)
    footer.add_run("\t")
    footer.add_run(website)
    for run in footer.runs:
        set_run_font(run, preset["font_east_asia"], preset["font_ascii"], 8, color="666666")
    add_paragraph_border(footer, "top", preset.get("accent_mid", preset["accent"]), size="4", space="4")


def format_list_paragraph(paragraph, preset: dict) -> None:
    paragraph.paragraph_format.left_indent = LIST_TEXT_INDENT
    paragraph.paragraph_format.first_line_indent = LIST_HANGING_INDENT
    paragraph.paragraph_format.line_spacing = preset["body_spacing"]
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(LIST_TEXT_INDENT, WD_TAB_ALIGNMENT.LEFT)


def iter_body_blocks(doc: Document):
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def is_blank_paragraph(block) -> bool:
    return isinstance(block, Paragraph) and not block.text.strip()


def apply_table_note_style(paragraph, preset: dict, normal_style) -> None:
    paragraph.style = normal_style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.page_break_before = False
    for run in paragraph.runs:
        set_run_font(run, preset["font_east_asia"], preset["font_ascii"], 9, False, "666666")


def move_table_notes_below_tables(doc: Document, preset: dict) -> None:
    normal_style = doc.styles["Normal"]
    blocks = list(iter_body_blocks(doc))
    for index, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        note_text = block.text.strip()
        if not TABLE_NOTE_RE.match(note_text):
            continue

        prev_nonblank = None
        for scan in range(index - 1, -1, -1):
            candidate = blocks[scan]
            if is_blank_paragraph(candidate):
                continue
            prev_nonblank = candidate
            break

        next_nonblank = None
        for scan in range(index + 1, len(blocks)):
            candidate = blocks[scan]
            if is_blank_paragraph(candidate):
                continue
            next_nonblank = candidate
            break

        if not isinstance(prev_nonblank, Table) and isinstance(next_nonblank, Table):
            next_nonblank._tbl.addnext(block._p)

        apply_table_note_style(block, preset, normal_style)


def paragraph_num_info(paragraph) -> tuple[int | None, int | None]:
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None:
        return None, None
    num_pr = p_pr.numPr
    num_id = num_pr.numId.val if num_pr.numId is not None else None
    ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
    return num_id, ilvl


def get_list_style_num_info(doc: Document, style_name: str = "List Number") -> tuple[int | None, int | None]:
    style = doc.styles[style_name]
    p_pr = style._element.pPr
    if p_pr is None or p_pr.numPr is None:
        return None, None
    num_pr = p_pr.numPr
    num_id = num_pr.numId.val if num_pr.numId is not None else None
    ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
    return num_id, ilvl


def allocate_numbering_clone(doc: Document, source_num_id: int) -> int:
    numbering = doc.part.numbering_part._element
    nums = numbering.findall(qn("w:num"))
    existing_ids = [int(item.get(qn("w:numId"))) for item in nums if item.get(qn("w:numId"))]
    source_num = None
    for item in nums:
        if item.get(qn("w:numId")) == str(source_num_id):
            source_num = item
            break
    if source_num is None:
        raise ValueError(f"Missing source numbering definition for numId={source_num_id}")
    new_num_id = (max(existing_ids) + 1) if existing_ids else (source_num_id + 1)
    new_num = copy.deepcopy(source_num)
    new_num.set(qn("w:numId"), str(new_num_id))
    numbering.append(new_num)
    return new_num_id


def set_paragraph_numbering(paragraph, num_id: int, ilvl: int = 0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl_node = num_pr.get_or_add_ilvl()
    ilvl_node.val = ilvl
    num_id_node = num_pr.get_or_add_numId()
    num_id_node.val = num_id


def is_numbering_restart_boundary(paragraph) -> bool:
    text = paragraph.text.strip()
    if not text:
        return False
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name.startswith("Heading"):
        return True
    if CN_SECTION_TITLE_RE.match(text):
        return True
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.bold for run in runs)


def reset_numbering_after_titles(doc: Document) -> None:
    base_num_id, base_ilvl = get_list_style_num_info(doc, "List Number")
    if base_num_id is None:
        return

    current_num_id = base_num_id
    restart_pending = True

    for paragraph in doc.paragraphs:
        if is_numbering_restart_boundary(paragraph):
            restart_pending = True
            continue

        style_name = paragraph.style.name if paragraph.style else ""
        if style_name != "List Number":
            continue

        if restart_pending:
            current_num_id = allocate_numbering_clone(doc, base_num_id)
            restart_pending = False

        _old_num_id, para_ilvl = paragraph_num_info(paragraph)
        set_paragraph_numbering(paragraph, current_num_id, para_ilvl if para_ilvl is not None else (base_ilvl or 0))


def polish_docx(input_path: Path, output_path: Path, style_name: str, title: str | None = None) -> None:
    preset = style_preset(style_name)
    doc = Document(input_path)

    for section_index, section in enumerate(doc.sections):
        if section.orientation == WD_ORIENTATION.LANDSCAPE or section.page_width > section.page_height:
            section.orientation = WD_ORIENTATION.LANDSCAPE
            section.page_width = Cm(29.7)
            section.page_height = Cm(21)
            section.top_margin = Cm(1.8)
            section.bottom_margin = Cm(1.8)
            section.left_margin = Cm(1.8)
            section.right_margin = Cm(1.8)
        else:
            section.orientation = WD_ORIENTATION.PORTRAIT
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(2.4)
            section.bottom_margin = Cm(2.2)
            section.left_margin = Cm(2.6)
            section.right_margin = Cm(2.4)
        apply_branded_page_furniture(section, preset, title or "", blank_first_page=section_index == 0)

    normal = doc.styles["Normal"]
    normal.font.name = preset["font_ascii"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), preset["font_east_asia"])
    normal.font.size = Pt(preset["body_size"])
    normal.paragraph_format.line_spacing = preset["body_spacing"]
    normal.paragraph_format.space_after = Pt(4)

    for level in range(1, 7):
        style = doc.styles[f"Heading {level}"]
        style.font.name = preset["heading_font_ascii"]
        style._element.rPr.rFonts.set(qn("w:eastAsia"), preset["heading_font_east_asia"])
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(preset["accent"] if level <= 2 else "333333")
        if level == 1:
            heading_size = float(preset["heading1"])
        else:
            heading_size = max(float(preset["heading2"]) - (level - 2) * 1.5, 11)
        style.font.size = Pt(heading_size)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(8 if level == 1 else 4)
        style.paragraph_format.page_break_before = False
        style.paragraph_format.keep_with_next = level >= 2

    for paragraph in doc.paragraphs:
        style_name_local = paragraph.style.name if paragraph.style else ""
        if style_name_local == "Normal" and paragraph.text.strip():
            paragraph.paragraph_format.line_spacing = preset["body_spacing"]
            paragraph.paragraph_format.space_after = Pt(4)
            if paragraph.paragraph_format.first_line_indent is None and paragraph.paragraph_format.left_indent in (None, 0):
                paragraph.paragraph_format.first_line_indent = BODY_FIRST_LINE_INDENT
        if style_name_local.startswith("List"):
            format_list_paragraph(paragraph, preset)
        if style_name_local.startswith("Heading"):
            level_text = style_name_local.replace("Heading", "").strip()
            level = int(level_text) if level_text.isdigit() else 1
            paragraph.paragraph_format.page_break_before = False
            paragraph.paragraph_format.keep_with_next = level >= 2
        if paragraph.text.strip().startswith(("图", "表")) or paragraph.style.name in {"Caption", "Figure"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = None
        for run in paragraph.runs:
            if style_name_local.startswith("Heading"):
                set_run_font(
                    run,
                    preset["heading_font_east_asia"],
                    preset["heading_font_ascii"],
                    None,
                    True,
                    preset["accent"] if style_name_local in {"Heading 1", "Heading 2"} else "333333",
                )
            else:
                set_run_font(run, preset["font_east_asia"], preset["font_ascii"])

    for table in doc.tables:
        try:
            table.style = "Table Grid"
        except Exception:
            pass
        if table.rows:
            repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                set_cell_margin(cell)
                if row_index == 0:
                    set_cell_shading(cell, preset["accent"])
                elif row_index % 2 == 1:
                    set_cell_shading(cell, preset["accent_light"])
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.first_line_indent = None
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            preset["font_east_asia"],
                            preset["font_ascii"],
                            preset["table_font"],
                            True if row_index == 0 else None,
                            "FFFFFF" if row_index == 0 else None,
                        )

    move_table_notes_below_tables(doc, preset)
    reset_numbering_after_titles(doc)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Polish an existing DOCX for bid delivery.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", "-o", type=Path, required=True)
    parser.add_argument("--style", default="formal-blue")
    parser.add_argument("--title", default=None)
    args = parser.parse_args()
    polish_docx(args.input.resolve(), args.output.resolve(), args.style, args.title)
    print(f"Wrote {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
