from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from audit_docx import audit_docx, audit_to_markdown
from build_docx_python import build_docx
from common import (
    PANDOC_ASSET_DIR,
    choose_reference_doc,
    extract_title,
    find_codex_runtime_tool,
    find_document_renderer,
    find_executable,
    infer_bid_title,
    read_text,
    run_command,
    skill_python,
    write_json,
    write_text,
)
from polish_docx import polish_docx
from preflight_markdown import analyze_markdown, report_to_markdown
from render_mermaid_blocks import render_blocks
from wps_pdf_export import export_pdf


STYLE_TO_THEME = {
    "hejia-yahei": "bid-blue",
    "formal-blue": "bid-blue",
    "compact-bid": "bid-blue",
    "government-red": "bid-red",
    "elegant-black": "bid-mono",
}


def next_versioned_docx_paths(base_stem: str, out_dir: Path) -> tuple[Path, Path]:
    stem = base_stem.strip() or "output"
    draft = out_dir / f"{stem}-draft.docx"
    base_final = out_dir / f"{stem}.docx"
    if not base_final.exists():
        return draft, base_final
    version = 2
    while True:
        candidate = out_dir / f"{stem} V{version}.docx"
        if not candidate.exists():
            return draft, candidate
        version += 1


def ensure_org_chart_present(docx_path: Path, figures_dir: Path) -> bool:
    chart_path = figures_dir / "fig-001-diagram.png"
    if not chart_path.is_file():
        return False
    doc = Document(docx_path)
    if len(doc.inline_shapes) >= 5:
        return False
    target_index = None
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if "图2-1" in text or "组织架构图" in text:
            target_index = index
            break
    if target_index is None:
        return False
    anchor = doc.paragraphs[target_index]._p
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = 1
    image_paragraph.paragraph_format.space_before = Pt(2)
    image_paragraph.paragraph_format.space_after = Pt(3)
    image_paragraph.add_run().add_picture(str(chart_path), width=Cm(15.6))
    anchor.addprevious(image_paragraph._p)
    doc.save(docx_path)
    return True


def build_with_pandoc(input_path: Path, output_path: Path, style: str, timeout: int) -> tuple[bool, str]:
    pandoc = find_executable("pandoc")
    if not pandoc:
        return False, "pandoc executable not found"
    reference = choose_reference_doc(style)
    lua_filter = PANDOC_ASSET_DIR / "markdown-to-docx.lua"
    if not reference or not reference.is_file():
        return False, "Pandoc reference DOCX template is missing"
    if not lua_filter.is_file():
        return False, "Pandoc Lua filter is missing"

    html_result = run_command([pandoc, str(input_path), "-f", "markdown", "-t", "html"], timeout=timeout, cwd=input_path.parent)
    if not html_result.ok:
        return False, html_result.stderr or html_result.stdout
    result = run_command(
        [
            pandoc,
            "-f",
            "html",
            "-o",
            str(output_path),
            "--reference-doc",
            str(reference),
            "--lua-filter",
            str(lua_filter),
            "--highlight-style",
            "tango",
        ],
        timeout=timeout,
        cwd=input_path.parent,
        stdin=html_result.stdout.encode("utf-8", errors="replace"),
    )
    if result.ok and output_path.is_file() and output_path.stat().st_size > 1000:
        return True, ""
    return False, result.stderr or result.stdout or "pandoc did not create DOCX"


def render_docx_if_available(docx_path: Path, out_dir: Path, timeout: int) -> dict:
    render_dir = out_dir / "rendered-pages"
    render_dir.mkdir(parents=True, exist_ok=True)
    render_script = find_document_renderer()
    if not render_script:
        return {"status": "skipped", "reason": "documents render_docx.py not found"}
    result = run_command([skill_python(), str(render_script), str(docx_path), "--output_dir", str(render_dir)], timeout=timeout)
    pngs = sorted(render_dir.glob("page-*.png"))
    if result.ok and pngs:
        return {"status": "rendered", "page_count": len(pngs), "output_dir": str(render_dir), "renderer": str(render_script)}
    return {
        "status": "failed",
        "reason": (result.stderr or result.stdout or "renderer did not produce page images")[:2000],
        "output_dir": str(render_dir),
        "renderer": str(render_script),
    }


def pdf_page_info(pdf_path: Path) -> dict:
    try:
        from pypdf import PdfReader
    except Exception:
        return {"page_count": None, "landscape_pages": []}
    try:
        reader = PdfReader(str(pdf_path))
    except Exception:
        return {"page_count": None, "landscape_pages": []}
    landscape_pages = []
    for index, page in enumerate(reader.pages, 1):
        box = page.mediabox
        if float(box.width) > float(box.height):
            landscape_pages.append(index)
    return {"page_count": len(reader.pages), "landscape_pages": landscape_pages}


def pick_pdf_sample_pages(page_count: int | None, landscape_pages: list[int], leading_pages: int = 5) -> list[int]:
    pages: set[int] = set()
    if page_count:
        pages.update(range(1, min(leading_pages, page_count) + 1))
        pages.add(page_count)
    else:
        pages.update(range(1, leading_pages + 1))
    pages.update(page for page in landscape_pages if page > 0 and (not page_count or page <= page_count))
    return sorted(pages)


def image_ink_ratio(image_path: Path) -> float | None:
    try:
        from PIL import Image
    except Exception:
        return None
    try:
        with Image.open(image_path) as image:
            gray = image.convert("L")
            width, height = gray.size
            pixels = gray.load()
            ink = 0
            total = width * height
            for y in range(height):
                for x in range(width):
                    if pixels[x, y] < 245:
                        ink += 1
            return ink / total if total else None
    except Exception:
        return None


def sparse_scan_pdf(pdf_path: Path, out_dir: Path, pdftoppm: Path, threshold: float = 0.006) -> dict:
    scan_dir = out_dir / "pdf-pages-fullscan"
    scan_dir.mkdir(parents=True, exist_ok=True)
    result = run_command([str(pdftoppm), "-png", "-r", "72", str(pdf_path), str(scan_dir / "page")], timeout=300)
    pages = sorted(scan_dir.glob("page-*.png"))
    if not result.ok or not pages:
        return {
            "status": "failed",
            "reason": (result.stderr or result.stdout or "full-page scan did not produce images")[:1200],
            "output_dir": str(scan_dir),
        }
    sparse_pages: list[dict] = []
    lowest: list[dict] = []
    for image_path in pages:
        ratio = image_ink_ratio(image_path)
        match = re.search(r"(\d+)", image_path.stem)
        page_number = int(match.group(1)) if match else None
        if ratio is None:
            continue
        item = {"page": page_number, "image": str(image_path), "ink_ratio": round(ratio, 5)}
        lowest.append(item)
        if ratio < threshold:
            sparse_pages.append(item)
    lowest = sorted(lowest, key=lambda item: item["ink_ratio"])[:10]
    return {
        "status": "scanned",
        "page_images": len(pages),
        "threshold": threshold,
        "sparse_pages": sparse_pages,
        "lowest_pages": lowest,
        "output_dir": str(scan_dir),
    }


def render_pdf_sample(pdf_path: Path, out_dir: Path, leading_pages: int = 5) -> dict:
    pdftoppm = find_codex_runtime_tool("pdftoppm")
    pdfinfo = find_codex_runtime_tool("pdfinfo")
    sample_dir = out_dir / "pdf-pages"
    sample_dir.mkdir(parents=True, exist_ok=True)
    page_data = pdf_page_info(pdf_path)
    page_count = page_data["page_count"]
    landscape_pages = page_data["landscape_pages"]
    if page_count is None and pdfinfo and pdfinfo.is_file():
        info = run_command([str(pdfinfo), str(pdf_path)], timeout=30)
        for line in info.stdout.splitlines():
            if line.startswith("Pages:"):
                try:
                    page_count = int(line.split(":", 1)[1].strip())
                except ValueError:
                    pass
    pages_to_render = pick_pdf_sample_pages(page_count, landscape_pages, leading_pages)
    if not pdftoppm or not pdftoppm.is_file():
        return {"status": "skipped", "reason": "pdftoppm not found", "pdf": str(pdf_path), "page_count": page_count}
    rendered: list[dict] = []
    failures: list[dict] = []
    for page_number in pages_to_render:
        prefix = sample_dir / f"page-{page_number:03d}"
        result = run_command(
            [str(pdftoppm), "-png", "-singlefile", "-r", "120", "-f", str(page_number), "-l", str(page_number), str(pdf_path), str(prefix)],
            timeout=120,
        )
        output = prefix.with_suffix(".png")
        if result.ok and output.is_file():
            ratio = image_ink_ratio(output)
            item = {"page": page_number, "image": str(output)}
            if ratio is not None:
                item["ink_ratio"] = round(ratio, 5)
                if ratio < 0.01:
                    item["sparse"] = True
            rendered.append(item)
        else:
            failures.append({"page": page_number, "reason": (result.stderr or result.stdout)[:1000]})
    sparse_scan = sparse_scan_pdf(pdf_path, out_dir, pdftoppm)
    if rendered and not failures:
        return {
            "status": "rendered",
            "pdf": str(pdf_path),
            "page_count": page_count,
            "sample_pages": len(rendered),
            "pages": rendered,
            "landscape_pages": landscape_pages,
            "output_dir": str(sample_dir),
            "renderer": str(pdftoppm),
            "sparse_scan": sparse_scan,
        }
    return {
        "status": "failed" if not rendered else "warn",
        "pdf": str(pdf_path),
        "page_count": page_count,
        "sample_pages": len(rendered),
        "pages": rendered,
        "landscape_pages": landscape_pages,
        "failures": failures,
        "output_dir": str(sample_dir),
        "renderer": str(pdftoppm),
        "sparse_scan": sparse_scan,
    }


def image_shape_warnings(diagrams: dict, ratio_limit: float = 3.2) -> list[dict]:
    warnings: list[dict] = []
    try:
        from PIL import Image
    except Exception:
        return warnings
    for item in diagrams.get("diagrams", []):
        image_path = Path(item.get("image", ""))
        if not image_path.is_file():
            continue
        with Image.open(image_path) as image:
            width, height = image.size
        ratio = width / height if height else 0
        if ratio > ratio_limit:
            warnings.append(
                {
                    "index": item.get("index"),
                    "caption": item.get("caption", ""),
                    "image": str(image_path),
                    "width": width,
                    "height": height,
                    "ratio": round(ratio, 2),
                }
            )
    return warnings


def wide_diagrams_accounted_for(wide_diagrams: list[dict], structural: dict, render: dict) -> bool:
    if not wide_diagrams:
        return True
    if structural.get("landscape_section_count", 0) < len(wide_diagrams):
        return False
    pdf_render = render.get("pdf_render", {}) if isinstance(render, dict) else {}
    if pdf_render and pdf_render.get("status") == "rendered":
        landscape_pages = pdf_render.get("landscape_pages", [])
        return len(landscape_pages) >= len(wide_diagrams)
    if render.get("status") == "rendered":
        return True
    return False


def sparse_pdf_pages(render: dict) -> list[dict]:
    pdf_render = render.get("pdf_render", {}) if isinstance(render, dict) else {}
    if not pdf_render:
        return []
    sparse = [item for item in pdf_render.get("pages", []) if item.get("sparse")]
    scan = pdf_render.get("sparse_scan", {})
    sparse.extend(scan.get("sparse_pages", []))
    seen: set[int] = set()
    unique = []
    for item in sparse:
        page = item.get("page")
        if page in seen:
            continue
        seen.add(page)
        unique.append(item)
    return unique


def qa_summary(preflight: dict, diagrams: dict, docx_path: Path, render: dict, structural: dict) -> dict:
    checks = []
    checks.append(
        {
            "name": "preflight_no_fatal",
            "status": "pass" if preflight["fatal_count"] == 0 else "fail",
            "detail": f"{preflight['fatal_count']} fatal issues",
        }
    )
    checks.append(
        {
            "name": "diagrams_rendered",
            "status": "pass" if diagrams["failed_count"] == 0 else "warn",
            "detail": f"{diagrams['diagram_count'] - diagrams['failed_count']}/{diagrams['diagram_count']} rendered",
        }
    )
    wide_diagrams = image_shape_warnings(diagrams)
    wide_ok = wide_diagrams_accounted_for(wide_diagrams, structural, render)
    checks.append(
        {
            "name": "wide_diagram_review",
            "status": "pass" if wide_ok else "warn",
            "detail": f"{len(wide_diagrams)} very wide diagram(s); landscape sections={structural.get('landscape_section_count', 0)}",
            "items": wide_diagrams,
        }
    )
    checks.append(
        {
            "name": "placeholders_resolved",
            "status": "pass" if preflight["placeholder_count"] == 0 and structural["placeholder_count"] == 0 else "warn",
            "detail": f"source={preflight['placeholder_count']}, docx={structural['placeholder_count']}",
        }
    )
    checks.append(
        {
            "name": "docx_exists",
            "status": "pass" if docx_path.is_file() and docx_path.stat().st_size > 1000 else "fail",
            "detail": str(docx_path),
        }
    )
    checks.append(
        {
            "name": "docx_structural_audit",
            "status": "pass" if structural["fatal_count"] == 0 else "fail",
            "detail": f"paragraphs={structural['paragraph_count']}, tables={structural['table_count']}, images={structural['image_count']}, sections={structural.get('section_count', 0)}, landscape={structural.get('landscape_section_count', 0)}, internal_prompts={structural.get('internal_prompt_count', 0)}, fatal={structural['fatal_count']}, warnings={structural['warning_count']}",
        }
    )
    sparse_pages = sparse_pdf_pages(render)
    checks.append(
        {
            "name": "pdf_sparse_page_review",
            "status": "pass" if not sparse_pages else "warn",
            "detail": f"{len(sparse_pages)} sparse sampled page(s)",
            "items": sparse_pages,
        }
    )
    render_status = render.get("status")
    checks.append(
        {
            "name": "render_review",
            "status": "pass" if render_status == "rendered" else "warn",
            "detail": json.dumps(render, ensure_ascii=False),
        }
    )
    return {
        "checks": checks,
        "pass_count": sum(1 for item in checks if item["status"] == "pass"),
        "warn_count": sum(1 for item in checks if item["status"] == "warn"),
        "fail_count": sum(1 for item in checks if item["status"] == "fail"),
    }


def qa_to_markdown(data: dict) -> str:
    lines = ["# DOCX Build QA", "", "| Check | Status | Detail |", "| --- | --- | --- |"]
    for item in data["checks"]:
        detail = item["detail"].replace("|", "\\|")
        lines.append(f"| {item['name']} | {item['status']} | {detail} |")
    wide_items = []
    for item in data["checks"]:
        if item["name"] == "wide_diagram_review":
            wide_items = item.get("items", [])
            break
    if wide_items:
        lines.extend(["", "## Wide Diagram Review", "", "| # | Ratio | Size | Image | Caption |", "| ---: | ---: | --- | --- | --- |"])
        for item in wide_items:
            caption = str(item.get("caption", "")).replace("|", "\\|")
            lines.append(
                f"| {item.get('index')} | {item.get('ratio')} | {item.get('width')}x{item.get('height')} | `{item.get('image')}` | {caption} |"
            )
    sparse_items = []
    for item in data["checks"]:
        if item["name"] == "pdf_sparse_page_review":
            sparse_items = item.get("items", [])
            break
    if sparse_items:
        lines.extend(["", "## Sparse PDF Page Review", "", "| Page | Ink Ratio | Image |", "| ---: | ---: | --- |"])
        for item in sparse_items:
            lines.append(f"| {item.get('page')} | {item.get('ink_ratio')} | `{item.get('image')}` |")
    lines.extend(
        [
            "",
            f"- Passed: {data['pass_count']}",
            f"- Warnings: {data['warn_count']}",
            f"- Failed: {data['fail_count']}",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert a large Markdown bid manuscript into polished DOCX.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--style", choices=["hejia-yahei", "formal-blue", "government-red", "elegant-black", "compact-bid"], default="formal-blue")
    parser.add_argument("--engine", choices=["auto", "pandoc", "python"], default="auto")
    parser.add_argument("--title", default=None)
    parser.add_argument("--output-base-name", default=None)
    parser.add_argument("--no-kroki", action="store_true")
    parser.add_argument("--skip-render", action="store_true")
    parser.add_argument("--timeout", type=int, default=120)
    args = parser.parse_args()

    input_path = args.input.resolve()
    if not input_path.is_file():
        parser.error(f"Input Markdown not found: {input_path}")
    out_dir = args.out_dir.resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    source_copy = out_dir / input_path.name
    if source_copy != input_path:
        shutil.copy2(input_path, source_copy)

    text = read_text(input_path)
    if args.title:
        doc_title = args.title
    else:
        doc_title = infer_bid_title(text, extract_title(text, input_path.stem))

    preflight = analyze_markdown(input_path)
    write_json(out_dir / "preflight-report.json", preflight)
    write_text(out_dir / "preflight-report.md", report_to_markdown(preflight))

    theme = STYLE_TO_THEME.get(args.style, "bid-blue")
    diagrams = render_blocks(input_path, out_dir, theme, not args.no_kroki, args.timeout)
    normalized = Path(diagrams["normalized_markdown"])

    output_base_name = args.output_base_name or input_path.stem
    draft_docx, final_docx = next_versioned_docx_paths(output_base_name, out_dir)
    engine_used = "python"
    build_error = ""

    if args.engine in {"auto", "pandoc"}:
        ok, build_error = build_with_pandoc(normalized, draft_docx, args.style, args.timeout)
        if ok:
            engine_used = "pandoc"
        elif args.engine == "pandoc":
            write_json(out_dir / "build-report.json", {"engine": "pandoc", "status": "failed", "error": build_error})
            print(f"Pandoc build failed: {build_error}")
            return 4

    if draft_docx.exists():
        draft_docx.unlink()
    build_docx(normalized, draft_docx, args.style, doc_title, extra_asset_dirs=[input_path.parent])
    engine_used = "python"

    polish_docx(draft_docx, final_docx, args.style, doc_title)
    ensure_org_chart_present(final_docx, out_dir / "figures")
    structural = audit_docx(final_docx)
    write_json(out_dir / "docx-audit.json", structural)
    write_text(out_dir / "docx-audit.md", audit_to_markdown(structural))
    render = {"status": "skipped", "reason": "--skip-render"}
    if not args.skip_render:
        render = render_docx_if_available(final_docx, out_dir, args.timeout)
        if render.get("status") != "rendered":
            pywin32_root = Path.cwd() / ".tmp_pydeps"
            pdf_path = out_dir / "final-wps.pdf"
            pdf_result = export_pdf(final_docx, pdf_path, pywin32_root if pywin32_root.is_dir() else None)
            if pdf_result.get("status") == "exported":
                pdf_render = render_pdf_sample(pdf_path, out_dir)
                render = {"status": "rendered" if pdf_render.get("status") == "rendered" else "warn", "backend": "wps-pdf", "export": pdf_result, "pdf_render": pdf_render}
            else:
                render = {"status": "failed", "backend": "docx-render-and-wps-pdf", "docx_render": render, "wps_export": pdf_result}

    build_report = {
        "input": str(input_path),
        "source_copy": str(source_copy),
        "style": args.style,
        "diagram_theme": theme,
        "engine": engine_used,
        "pandoc_error": build_error,
        "draft_docx": str(draft_docx),
        "final_docx": str(final_docx),
        "render": render,
    }
    write_json(out_dir / "build-report.json", build_report)
    qa = qa_summary(preflight, diagrams, final_docx, render, structural)
    write_json(out_dir / "qa-report.json", qa)
    write_text(out_dir / "qa-report.md", qa_to_markdown(qa))

    print(f"Final DOCX: {final_docx}")
    print(f"Engine: {engine_used}")
    print(f"QA: {out_dir / 'qa-report.md'}")
    return 0 if qa["fail_count"] == 0 else 5


if __name__ == "__main__":
    raise SystemExit(main())
