from __future__ import annotations

import argparse
import html
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENTATION, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html as lxml_html
from PIL import Image

from common import (
    DEFAULT_COMPANY_NAME,
    DEFAULT_COMPANY_SLOGAN,
    DEFAULT_COMPANY_WEBSITE,
    HEADING_RE,
    infer_bid_title,
    read_text,
    strip_inline_markup,
    style_preset,
)


LIST_RE = re.compile(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$")
STYLE_BLOCK_RE = re.compile(r"<style\b.*?</style>", re.IGNORECASE | re.DOTALL)
IMG_SRC_RE = re.compile(r"<img\b|!\[[^\]]*\]\([^)]+\)", re.IGNORECASE)
TABLE_START_RE = re.compile(r"<table\b", re.IGNORECASE)
TABLE_END_RE = re.compile(r"</table>", re.IGNORECASE)
PART_PAGE_RE = re.compile(r'<div\s+class="([^"]*\bpart-page\b[^"]*)"[^>]*>(.*?)</div>', re.IGNORECASE | re.DOTALL)
PAGE_BREAK_RE = re.compile(
    r'(class="[^"]*\bpage-break\b[^"]*"|<br\s+clear="all"\s*/?>|<span\s+class="page-break"\s*>\s*</span>)',
    re.IGNORECASE,
)
BODY_FIRST_LINE_INDENT = Cm(0.74)
TABLE_NOTE_RE = re.compile(r"^(表\d+(?:[-－—]\d+[A-Za-z]?)?\s+.+)$")
LANDSCAPE_MIN_COLS = 9
LANDSCAPE_FORCE_MIN_CELL_TEXT = 18


@dataclass
class BuildStats:
    html_images_seen: int = 0
    images_inserted: int = 0
    html_tables: int = 0
    markdown_tables: int = 0
    page_breaks: int = 0
    part_pages: int = 0
    underlined_runs: int = 0
    landscape_tables: int = 0


def set_run_font(
    run,
    preset: dict,
    *,
    east_asia: str | None = None,
    ascii_font: str | None = None,
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
    underline: bool | None = None,
    italic: bool | None = None,
) -> None:
    ea = east_asia or preset["font_east_asia"]
    ascii_name = ascii_font or preset["font_ascii"]
    run.font.name = ascii_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), ea)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if underline is not None:
        run.underline = underline
    if italic is not None:
        run.italic = italic


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    run._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    paragraph.add_run("")
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_paragraph_border(paragraph, edge: str, color: str, size: str = "6", space: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = p_bdr.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        p_bdr.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)


def clear_story(story):
    for child in list(story._element):
        story._element.remove(child)
    return story.add_paragraph()


def reset_paragraph(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = None
    fmt.left_indent = None
    fmt.right_indent = None
    fmt.line_spacing = 1
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def section_text_width(section):
    return section.page_width - section.left_margin - section.right_margin


def apply_page_furniture(section, preset: dict, doc_title: str, *, blank_first_page: bool = False) -> None:
    section.header_distance = Cm(0.85)
    section.footer_distance = Cm(0.9)
    section.different_first_page_header_footer = blank_first_page
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    if blank_first_page:
        clear_story(section.first_page_header)
        clear_story(section.first_page_footer)

    width = section_text_width(section)

    header = clear_story(section.header)
    reset_paragraph(header)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    left = header.add_run(DEFAULT_COMPANY_NAME)
    set_run_font(
        left,
        preset,
        east_asia=preset["heading_font_east_asia"],
        ascii_font=preset["heading_font_ascii"],
        size=8.5,
        bold=True,
        color=preset["accent"],
    )
    header.add_run("\t")
    right = header.add_run(doc_title)
    set_run_font(right, preset, size=8, color="666666")
    add_paragraph_border(header, "bottom", preset["accent"], size="6", space="2")

    footer = clear_story(section.footer)
    reset_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.paragraph_format.tab_stops.add_tab_stop(width // 2, WD_TAB_ALIGNMENT.CENTER)
    footer.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run(" / ".join(DEFAULT_COMPANY_SLOGAN.split()))
    footer.add_run("\t")
    footer.add_run("第")
    add_field(footer, "PAGE")
    footer.add_run("页")
    footer.add_run("\t")
    footer.add_run(DEFAULT_COMPANY_WEBSITE)
    for run in footer.runs:
        set_run_font(run, preset, size=8, color="666666")
    add_paragraph_border(footer, "top", preset.get("accent_mid", preset["accent"]), size="4", space="4")


def configure_document(doc: Document, preset: dict, doc_title: str) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENTATION.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    apply_page_furniture(section, preset, doc_title, blank_first_page=True)

    normal = doc.styles["Normal"]
    normal.font.name = preset["font_ascii"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), preset["font_east_asia"])
    normal.font.size = Pt(preset["body_size"])
    normal.paragraph_format.line_spacing = preset["body_spacing"]
    normal.paragraph_format.space_after = Pt(4)

    for level in range(1, 7):
        style = doc.styles[f"Heading {level}"]
        style.font.name = preset["heading_font_ascii"]
        style._element.rPr.rFonts.set(qn("w:eastAsia"), preset["heading_font_east_asia"])
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(preset["accent"] if level <= 2 else "333333")
        style.font.size = Pt(max(float(preset["heading1"]) - (level - 1) * 1.35, 10.5))
        style.paragraph_format.space_before = Pt(14 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(8 if level == 1 else 4)
        style.paragraph_format.keep_with_next = level >= 2


def set_portrait(section, preset: dict, doc_title: str) -> None:
    section.orientation = WD_ORIENTATION.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    apply_page_furniture(section, preset, doc_title)


def set_landscape(section, preset: dict, doc_title: str) -> None:
    section.orientation = WD_ORIENTATION.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    apply_page_furniture(section, preset, doc_title)


def should_use_landscape(col_count: int, rows: list[list[str]] | None = None) -> bool:
    if col_count >= LANDSCAPE_MIN_COLS:
        return True
    if col_count < 8 or not rows:
        return False
    wide_cells = 0
    inspect_rows = rows[: min(len(rows), 10)]
    for row in inspect_rows:
        for cell in row:
            if len(strip_inline_markup(cell).strip()) >= LANDSCAPE_FORCE_MIN_CELL_TEXT:
                wide_cells += 1
    return wide_cells >= max(6, col_count)


def begin_landscape_if_needed(doc: Document, col_count: int, preset: dict, doc_title: str, stats: BuildStats, rows: list[list[str]] | None = None) -> bool:
    if not should_use_landscape(col_count, rows):
        return False
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(section, preset, doc_title)
    stats.landscape_tables += 1
    return True


def end_landscape_if_needed(doc: Document, active: bool, preset: dict, doc_title: str) -> None:
    if not active:
        return
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_portrait(section, preset, doc_title)


def add_cover(doc: Document, preset: dict, doc_title: str) -> None:
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(doc_title)
    set_run_font(
        run,
        preset,
        east_asia=preset["heading_font_east_asia"],
        ascii_font=preset["heading_font_ascii"],
        size=24,
        bold=True,
        color=preset["accent"],
    )
    p.paragraph_format.space_after = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("资信技术标")
    set_run_font(
        run,
        preset,
        east_asia=preset["heading_font_east_asia"],
        ascii_font=preset["heading_font_ascii"],
        size=22,
        bold=True,
        color="333333",
    )
    p.paragraph_format.space_after = Pt(18)

    rule = doc.add_paragraph()
    rule.paragraph_format.left_indent = Cm(4.0)
    rule.paragraph_format.right_indent = Cm(4.0)
    add_paragraph_border(rule, "bottom", preset["accent"], size="12", space="4")

    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"投标人：{DEFAULT_COMPANY_NAME}")
    set_run_font(run, preset, size=14, bold=True, color="333333")
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026年6月")
    set_run_font(run, preset, size=12, color="333333")
    doc.add_page_break()


def collect_toc_entries(text: str, max_level: int = 2) -> list[tuple[int, str]]:
    entries: list[tuple[int, str]] = []
    in_style = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.lower().startswith("<style"):
            in_style = True
        if in_style:
            if "</style>" in stripped.lower():
                in_style = False
            continue
        match = HEADING_RE.match(stripped)
        if match and len(match.group(1)) <= max_level:
            title = strip_inline_markup(match.group(2)).strip()
            if title:
                entries.append((len(match.group(1)), title))
    return entries


def add_static_toc(doc: Document, preset: dict, entries: list[tuple[int, str]]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目录")
    set_run_font(
        run,
        preset,
        east_asia=preset["heading_font_east_asia"],
        ascii_font=preset["heading_font_ascii"],
        size=18,
        bold=True,
        color=preset["accent"],
    )
    p.paragraph_format.space_after = Pt(12)
    for level, title in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8 if level == 2 else 0)
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(title)
        set_run_font(run, preset, size=9 if level == 2 else 10, bold=level == 1, color=preset["accent"] if level == 1 else "333333")
    doc.add_page_break()


def add_part_page(doc: Document, preset: dict, title: str, stats: BuildStats) -> None:
    stats.part_pages += 1
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title.strip())
    set_run_font(
        run,
        preset,
        east_asia=preset["heading_font_east_asia"],
        ascii_font=preset["heading_font_ascii"],
        size=30,
        bold=True,
        color=preset["accent"],
    )
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(80)
    doc.add_page_break()
    stats.page_breaks += 1


def normalize_text(value: str) -> str:
    value = html.unescape(value or "")
    value = value.replace("\xa0", " ")
    return value


def parse_style_value(style: str, key: str) -> str | None:
    for item in style.split(";"):
        if ":" not in item:
            continue
        left, right = item.split(":", 1)
        if left.strip().lower() == key:
            return right.strip()
    return None


def style_size(style: str, default: float | None = None) -> float | None:
    raw = parse_style_value(style, "font-size")
    if not raw:
        return default
    match = re.search(r"([\d.]+)\s*pt", raw)
    return float(match.group(1)) if match else default


def style_color(style: str, default: str | None = None) -> str | None:
    raw = parse_style_value(style, "color")
    if not raw:
        return default
    match = re.search(r"#?([0-9a-fA-F]{6})", raw)
    return match.group(1).upper() if match else default


def html_text(line_or_node) -> str:
    try:
        if isinstance(line_or_node, str):
            node = lxml_html.fragment_fromstring(line_or_node, create_parent="div")
        else:
            node = line_or_node
        return normalize_text("".join(node.itertext())).strip()
    except Exception:
        return normalize_text(re.sub(r"<[^>]+>", "", str(line_or_node))).strip()


def format_body_paragraph(paragraph, preset: dict) -> None:
    paragraph.paragraph_format.first_line_indent = BODY_FIRST_LINE_INDENT
    paragraph.paragraph_format.line_spacing = preset["body_spacing"]
    paragraph.paragraph_format.space_after = Pt(4)


def add_text_run(paragraph, text: str, state: dict, preset: dict, stats: BuildStats) -> None:
    text = normalize_text(text)
    if text == "":
        return
    run = paragraph.add_run(text)
    if state.get("underline"):
        stats.underlined_runs += 1
    set_run_font(
        run,
        preset,
        size=state.get("size", preset["body_size"]),
        bold=state.get("bold"),
        color=state.get("color"),
        underline=state.get("underline"),
        italic=state.get("italic"),
    )


def add_html_inline(paragraph, node, state: dict, preset: dict, stats: BuildStats) -> None:
    if node.text:
        add_text_run(paragraph, node.text, state, preset, stats)
    for child in node:
        child_state = dict(state)
        tag = child.tag.lower() if isinstance(child.tag, str) else ""
        child_style = child.get("style", "")
        if tag in {"strong", "b"}:
            child_state["bold"] = True
        if tag in {"em", "i"}:
            child_state["italic"] = True
        if tag == "u":
            child_state["underline"] = True
        if child_style:
            child_state["size"] = style_size(child_style, child_state.get("size"))
            child_state["color"] = style_color(child_style, child_state.get("color"))
        if tag == "br":
            paragraph.add_run().add_break()
        else:
            add_html_inline(paragraph, child, child_state, preset, stats)
        if child.tail:
            add_text_run(paragraph, child.tail, state, preset, stats)


def add_html_paragraph(doc: Document, line: str, preset: dict, stats: BuildStats) -> None:
    node = lxml_html.fragment_fromstring(line, create_parent="div")
    p_node = node.find(".//p")
    if p_node is None:
        text = html_text(node)
        if not text:
            return
        p = doc.add_paragraph()
        format_body_paragraph(p, preset)
        run = p.add_run(text)
        set_run_font(run, preset, size=preset["body_size"])
        return

    text = html_text(p_node)
    if not text:
        return
    p = doc.add_paragraph()
    style = p_node.get("style", "")
    align = (parse_style_value(style, "text-align") or "").lower()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = None
    else:
        format_body_paragraph(p, preset)
        indent = parse_style_value(style, "text-indent")
        if indent:
            match = re.search(r"([\d.]+)\s*pt", indent)
            if match:
                p.paragraph_format.first_line_indent = Pt(float(match.group(1)))
    add_html_inline(p, p_node, {"size": style_size(style, preset["body_size"]), "color": style_color(style)}, preset, stats)


def is_page_break_line(line: str) -> bool:
    stripped = line.strip().lower()
    return bool(PAGE_BREAK_RE.search(stripped)) and "img" not in stripped


def add_page_break(doc: Document, stats: BuildStats) -> None:
    doc.add_page_break()
    stats.page_breaks += 1


def extract_image_info(line: str) -> tuple[str, str, float | None]:
    if line.strip().startswith("!["):
        match = re.search(r'!\[([^\]]*)\]\(([^)\s]+)(?:\s+"([^"]*)")?\)', line)
        if not match:
            raise ValueError("no markdown image found")
        alt, src, title = match.groups()
        return src, title or alt or "", None

    node = lxml_html.fragment_fromstring(line, create_parent="div")
    img = node.find(".//img")
    if img is None:
        raise ValueError("no img tag found")
    src = img.get("src") or ""
    alt = img.get("alt") or ""
    css_width = None
    style = img.get("style") or ""
    match = re.search(r"width:\s*([\d.]+)\s*cm", style, re.IGNORECASE)
    if match:
        css_width = float(match.group(1))
    return src, alt, css_width


def caption_candidate(line: str) -> str | None:
    lowered = line.lower()
    if (("<p" not in lowered and not line.strip().startswith("![")) or "<img" in lowered or is_page_break_line(line)):
        return None
    text = html_text(line)
    if not text:
        return None
    return text if len(text) <= 120 else None


def resolve_resource_path(src: str, base_dir: Path, extra_dirs: list[Path] | None = None) -> Path:
    extra_dirs = extra_dirs or []
    candidates = [(base_dir / src).resolve()]
    for root in extra_dirs:
        candidates.append((root / src).resolve())
    candidates.append((Path.cwd() / src).resolve())
    seen: set[str] = set()
    for candidate in candidates:
        key = str(candidate)
        if key in seen:
            continue
        seen.add(key)
        if candidate.is_file():
            return candidate
    return candidates[0]


def add_image(
    doc: Document,
    base_dir: Path,
    extra_dirs: list[Path],
    src: str,
    alt: str,
    caption: str | None,
    css_width: float | None,
    preset: dict,
    stats: BuildStats,
) -> None:
    stats.html_images_seen += 1
    image_path = resolve_resource_path(src, base_dir, extra_dirs)
    if not image_path.is_file():
        p = doc.add_paragraph()
        format_body_paragraph(p, preset)
        run = p.add_run(f"[图片缺失：{src}]")
        set_run_font(run, preset, size=preset["body_size"], color="9E1B1B")
        return

    is_contract = "合同扫描件" in src or ("合同-" in (caption or alt))
    try:
        with Image.open(image_path) as image:
            width_px, height_px = image.size
    except Exception:
        width_px, height_px = (1000, 1400)

    aspect = height_px / width_px if width_px else 1.4
    max_width = 15.6
    if css_width:
        max_width = min(max_width, css_width)
    if is_contract:
        max_width = min(max_width, 14.6)
    max_height = 21.4 if is_contract else 21.8
    width_cm = min(max_width, max_height / aspect if aspect else max_width)
    if width_cm < 6:
        width_cm = min(max_width, 6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    stats.images_inserted += 1

    label = caption or alt
    if label:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.first_line_indent = None
        cp.paragraph_format.space_after = Pt(5)
        run = cp.add_run(label)
        set_run_font(run, preset, size=10.5 if is_contract else 9, bold=is_contract, color="333333" if is_contract else "666666")


def set_cell_margin(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def style_table(table, preset: dict, *, header: bool = True, compact: bool = True) -> None:
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    table.autofit = True
    if table.rows:
        repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margin(cell)
            if header and row_index == 0:
                set_cell_shading(cell, preset["accent"])
            elif row_index % 2 == 1:
                set_cell_shading(cell, preset["accent_light"])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.first_line_indent = None
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        preset,
                        size=preset["table_font"] if compact else max(float(preset["table_font"]), 10),
                        bold=True if header and row_index == 0 else None,
                        color="FFFFFF" if header and row_index == 0 else None,
                    )


def add_html_table(doc: Document, html_block: str, preset: dict, doc_title: str, stats: BuildStats) -> None:
    stats.html_tables += 1
    root = lxml_html.fragment_fromstring(html_block, create_parent="div")
    table_node = root.find(".//table")
    if table_node is None:
        return

    parsed_rows: list[list[dict]] = []
    occupied: dict[tuple[int, int], bool] = {}
    max_cols = 0
    trs = table_node.xpath(".//tr")
    for r_index, tr in enumerate(trs):
        cells: list[dict] = []
        c_index = 0
        for cell_node in tr.xpath("./th|./td"):
            while occupied.get((r_index, c_index)):
                c_index += 1
            colspan = int(cell_node.get("colspan") or "1")
            rowspan = int(cell_node.get("rowspan") or "1")
            cell = {"row": r_index, "col": c_index, "rowspan": rowspan, "colspan": colspan, "node": cell_node}
            cells.append(cell)
            for rr in range(r_index, r_index + rowspan):
                for cc in range(c_index, c_index + colspan):
                    occupied[(rr, cc)] = True
            c_index += colspan
        max_cols = max(max_cols, c_index)
        parsed_rows.append(cells)

    if not parsed_rows or max_cols == 0:
        return

    html_rows_for_width = []
    for row_cells in parsed_rows:
        row_text = [html_text(item["node"]) for item in row_cells]
        html_rows_for_width.append(row_text)
    landscape = begin_landscape_if_needed(doc, max_cols, preset, doc_title, stats, html_rows_for_width)
    table = doc.add_table(rows=len(parsed_rows), cols=max_cols)
    for row_cells in parsed_rows:
        for item in row_cells:
            start = table.cell(item["row"], item["col"])
            target = table.cell(item["row"] + item["rowspan"] - 1, item["col"] + item["colspan"] - 1)
            if start is not target:
                start = start.merge(target)
            start.text = ""
            p = start.paragraphs[0]
            p.paragraph_format.first_line_indent = None
            add_html_inline(p, item["node"], {"size": preset["table_font"]}, preset, stats)
    style_table(table, preset, header=True, compact=False)
    end_landscape_if_needed(doc, landscape, preset, doc_title)


def split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    cells: list[str] = []
    current: list[str] = []
    escaped = False
    for char in stripped:
        if escaped:
            current.append(char)
            escaped = False
        elif char == "\\":
            escaped = True
        elif char == "|":
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    cells.append("".join(current).strip())
    return cells


def clean_markdown_table_cell(text: str) -> str:
    text = strip_inline_markup(text)
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    return text.strip()


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return len(cells) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def next_nonempty_line_index(lines: list[str], start: int) -> int | None:
    for index in range(start, len(lines)):
        if lines[index].strip():
            return index
    return None


def table_note_targets_following_table(lines: list[str], note_index: int) -> bool:
    header_index = next_nonempty_line_index(lines, note_index + 1)
    if header_index is None:
        return False
    header_line = lines[header_index].strip()
    if TABLE_START_RE.search(header_line):
        return True
    if "|" not in header_line:
        return False
    separator_index = next_nonempty_line_index(lines, header_index + 1)
    if separator_index is None:
        return False
    return is_table_separator(lines[separator_index].strip())


def add_markdown_table(doc: Document, header: list[str], rows: list[list[str]], preset: dict, doc_title: str, stats: BuildStats) -> None:
    stats.markdown_tables += 1
    col_count = max([len(header), *(len(row) for row in rows)] or [len(header)])
    landscape = begin_landscape_if_needed(doc, col_count, preset, doc_title, stats, [header, *rows])
    table = doc.add_table(rows=1, cols=col_count)
    for c in range(col_count):
        table.cell(0, c).text = clean_markdown_table_cell(header[c]) if c < len(header) else ""
    for row in rows:
        cells = table.add_row().cells
        for c in range(col_count):
            cells[c].text = clean_markdown_table_cell(row[c]) if c < len(row) else ""
    style_table(table, preset, header=True, compact=True)
    end_landscape_if_needed(doc, landscape, preset, doc_title)


def add_table_note(doc: Document, note_text: str, preset: dict) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(strip_inline_markup(note_text))
    set_run_font(run, preset, size=9, color="666666")


def add_markdown_inline(paragraph, text: str, preset: dict, stats: BuildStats) -> None:
    pos = 0
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|<u>.*?</u>)", re.IGNORECASE)
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(strip_inline_markup(text[pos : match.start()]))
            set_run_font(run, preset, size=preset["body_size"])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token.strip("*"))
            set_run_font(run, preset, size=preset["body_size"], bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token.strip("`"))
            set_run_font(run, preset, east_asia="Consolas", ascii_font="Consolas", size=8.5, color=preset["accent"])
        else:
            stats.underlined_runs += 1
            run = paragraph.add_run(html_text(token))
            set_run_font(run, preset, size=preset["body_size"], underline=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(strip_inline_markup(text[pos:]))
        set_run_font(run, preset, size=preset["body_size"])


def add_heading(doc: Document, level: int, title: str, preset: dict) -> None:
    p = doc.add_heading(strip_inline_markup(title), level=level)
    if level == 1:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = level >= 2
    for run in p.runs:
        set_run_font(
            run,
            preset,
            east_asia=preset["heading_font_east_asia"],
            ascii_font=preset["heading_font_ascii"],
            size=max(float(preset["heading1"]) - (level - 1) * 1.35, 10.5),
            bold=True,
            color=preset["accent"] if level <= 2 else "333333",
        )


def parse_markdown(
    text: str,
    doc: Document,
    base_dir: Path,
    preset: dict,
    doc_title: str,
    extra_asset_dirs: list[Path] | None = None,
) -> BuildStats:
    stats = BuildStats()
    extra_asset_dirs = extra_asset_dirs or []
    text = STYLE_BLOCK_RE.sub("", text)
    lines = text.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    pending_table_note: str | None = None

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        table_note_match = TABLE_NOTE_RE.match(strip_inline_markup(stripped))
        if table_note_match and table_note_targets_following_table(lines, index):
            pending_table_note = table_note_match.group(1)
            index += 1
            continue

        part_match = PART_PAGE_RE.search(stripped)
        if part_match:
            if "part-page-before" in part_match.group(1):
                add_page_break(doc, stats)
            add_part_page(doc, preset, html_text(part_match.group(2)), stats)
            index += 1
            continue

        if is_page_break_line(stripped):
            add_page_break(doc, stats)
            index += 1
            continue

        if stripped.startswith("```") or stripped.startswith("~~~"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = None
                for code_line in code_lines:
                    run = p.add_run(code_line)
                    set_run_font(run, preset, east_asia="Consolas", ascii_font="Consolas", size=8.5, color="333333")
                    run.add_break()
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if TABLE_START_RE.search(stripped):
            table_lines = [line]
            while index < len(lines) and not TABLE_END_RE.search(lines[index]):
                index += 1
                if index < len(lines):
                    table_lines.append(lines[index])
            add_html_table(doc, "\n".join(table_lines), preset, doc_title, stats)
            if pending_table_note:
                add_table_note(doc, pending_table_note, preset)
                pending_table_note = None
            index += 1
            continue

        if IMG_SRC_RE.search(stripped):
            src, alt, css_width = extract_image_info(stripped)
            caption = None
            consume_until = index
            look = index + 1
            while look < len(lines) and not lines[look].strip():
                look += 1
            if look < len(lines):
                caption = caption_candidate(lines[look].strip())
                if caption is not None:
                    consume_until = look
            add_image(doc, base_dir, extra_asset_dirs, src, alt, caption, css_width, preset, stats)
            index = consume_until + 1
            continue

        if index < len(lines) - 1 and "|" in line and is_table_separator(lines[index + 1]):
            header = split_table_row(line)
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and "|" in lines[index].strip():
                rows.append(split_table_row(lines[index]))
                index += 1
            add_markdown_table(doc, header, rows, preset, doc_title, stats)
            if pending_table_note:
                add_table_note(doc, pending_table_note, preset)
                pending_table_note = None
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            add_heading(doc, min(len(heading.group(1)), 6), heading.group(2), preset)
            index += 1
            continue

        if stripped.lower().startswith("<p"):
            add_html_paragraph(doc, stripped, preset, stats)
            index += 1
            continue

        list_match = LIST_RE.match(line)
        if list_match:
            marker = list_match.group(2)
            content = list_match.group(3)
            style = "List Number" if marker[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.line_spacing = preset["body_spacing"]
            p.paragraph_format.space_after = Pt(2)
            add_markdown_inline(p, content, preset, stats)
            index += 1
            continue

        p = doc.add_paragraph()
        format_body_paragraph(p, preset)
        add_markdown_inline(p, stripped, preset, stats)
        pending_table_note = None
        index += 1

    return stats


def build_docx(
    input_path: Path,
    output_path: Path,
    style_name: str,
    title: str | None = None,
    extra_asset_dirs: list[Path] | None = None,
) -> None:
    text = read_text(input_path)
    doc_title = title or infer_bid_title(text, input_path.stem)
    preset = style_preset(style_name)

    doc = Document()
    configure_document(doc, preset, doc_title)
    add_cover(doc, preset, doc_title)
    add_static_toc(doc, preset, collect_toc_entries(text, max_level=2))
    parse_markdown(text, doc, input_path.parent, preset, doc_title, extra_asset_dirs)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build DOCX from Markdown using python-docx fallback.")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", "-o", type=Path, required=True)
    parser.add_argument("--style", default="formal-blue")
    parser.add_argument("--title", default=None)
    args = parser.parse_args()

    build_docx(args.input.resolve(), args.output.resolve(), args.style, args.title)
    print(f"Wrote {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
